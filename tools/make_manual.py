"""Generate the JazzWriter beginner's guide as a Word (.docx) document.

Run::

    python tools/make_manual.py

Output: ``docs/JazzWriter_新手引导.docx``
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "JazzWriter_新手引导.docx"
CJK = "微软雅黑"


def _cjk(font, name: str = CJK) -> None:
    font.name = name
    rpr = font._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), name)


def _setup(doc: Document) -> None:
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3",
                       "List Bullet", "List Number"):
        try:
            _cjk(doc.styles[style_name].font)
        except KeyError:
            pass
    doc.styles["Normal"].font.size = Pt(10.5)
    for name, size, color in (("Heading 1", 16, "1F4E79"), ("Heading 2", 13, "2E75B6"),
                              ("Heading 3", 11.5, "404040")):
        try:
            st = doc.styles[name].font
            st.size = Pt(size)
            st.color.rgb = RGBColor.from_string(color)
        except KeyError:
            pass


def h1(doc, text):
    doc.add_heading(text, level=1)


def h2(doc, text):
    doc.add_heading(text, level=2)


def p(doc, text, bold=False):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = bold
    return para


def bullets(doc, items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def numbers(doc, items):
    for it in items:
        doc.add_paragraph(it, style="List Number")


def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, htext in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = htext
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    doc.add_paragraph()
    return t


def build() -> Document:
    doc = Document()
    _setup(doc)

    # -- cover --------------------------------------------------------------
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    trun = title.add_run("JazzWriter 新手引导")
    trun.bold = True
    trun.font.size = Pt(28)
    _cjk(trun.font)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    srun = sub.add_run("长篇小说结构即兴引擎 · 从零开始的完整操作手册")
    srun.font.size = Pt(13)
    srun.font.color.rgb = RGBColor.from_string("666666")
    _cjk(srun.font)
    ver = doc.add_paragraph()
    ver.alignment = WD_ALIGN_PARAGRAPH.CENTER
    vrun = ver.add_run("适用版本 v2.18　·　Novel Jazz Engine")
    vrun.font.size = Pt(11)
    vrun.font.color.rgb = RGBColor.from_string("999999")
    _cjk(vrun.font)
    doc.add_paragraph()

    # -- TOC ----------------------------------------------------------------
    h1(doc, "目录")
    bullets(doc, [
        "一、这是什么（60 秒了解）",
        "二、安装与启动",
        "三、5 分钟快速上手",
        "四、界面速览",
        "五、核心概念：乐谱 · 结构 · 人物与关系",
        "六、结构工作流详解",
        "七、局部演练：正文 / 多版本 / Jam",
        "八、接入真实模型",
        "九、命令行速查",
        "十、常见问题（FAQ）",
        "十一、文件与目录说明",
        "十二、术语表",
    ])

    # -- 1 ----------------------------------------------------------------
    h1(doc, "一、这是什么（60 秒了解）")
    p(doc, "JazzWriter 不是“一键生成整本小说”的工具，而是一个帮你探索故事的“即兴引擎”。"
           "它的主要产物是：故事方案、全书梗概、分卷结构、逐章摘要卡片，以及人物与关系网络。"
           "正文生成是可选的“局部演练”。")
    p(doc, "工作循环一句话：", bold=True)
    p(doc, "给定一个前提 → 发散出多个实质不同的故事方向 → 采纳或融合 → 生成完整结构（梗概+分卷+逐章卡片）"
           "→ 校验 → 局部重新即兴（先出候选）→ 导出。最后由你判断与选择。")
    p(doc, "它适合：长篇小说/网文/系列的作者、需要管理人物与伏笔的写作者、想先想清楚“故事怎么成立”的人。")

    # -- 2 ----------------------------------------------------------------
    h1(doc, "二、安装与启动")
    p(doc, "1）安装依赖（在项目根目录执行）：")
    p(doc, "pip install -r requirements.txt")
    p(doc, "2）启动图形界面（推荐给新用户）：")
    p(doc, "python gui.py")
    p(doc, "或者直接双击根目录的 JazzWriter.bat。")
    p(doc, "3）也可以用命令行（CLI）：")
    p(doc, "python main.py --help")
    p(doc, "说明：默认使用离线 mock 模型，无需任何 API Key 即可跑通流程（但内容偏模板化）。"
           "接入真实模型见第八章。")

    # -- 3 ----------------------------------------------------------------
    h1(doc, "三、5 分钟快速上手")
    p(doc, "打开 GUI 后，按下面 6 步操作即可得到一套完整的故事结构。")

    h2(doc, "第 1 步：填写故事前提")
    p(doc, "在左侧“场景 / 章节摘要”框里写一句故事前提，例如：")
    p(doc, "一场技术事故改变了一个家庭。")
    p(doc, "并设置：即兴模式、全书章数、结构模板（如 mystery 悬疑）。")

    h2(doc, "第 2 步：生成创意方案")
    p(doc, "点「生成创意方案」→ 右侧「创意候选」标签页出现多个方向。"
           "每个方向在“因果机制 / 真相性质 / 人物关系 / 结局方向”上实质不同。")

    h2(doc, "第 3 步：采纳或融合")
    p(doc, "在“采纳方案#”里选一个编号；也可点「融合所选两个方案」把两个方向合并。"
           "然后点「采纳并展开全书结构」。")

    h2(doc, "第 4 步：查看与校验结构")
    p(doc, "右侧「全书结构 / 结构树 / 章节卡片 / 人物」展示分卷、逐章摘要与人物使用情况。"
           "点「校验结构」查看问题清单——**先列问题，再列参考分数**。")
    p(doc, "校验分两层：结构硬校验（缺章/编号/断链/未回收线索/闲置人物）+ 叙事审阅"
           "（事件是否合理、人物选择是否成立、结局是否回应开篇，给出依据与章号）。")
    p(doc, "另有一项显式检查：与“已确定事实”的语义冲突（如事实称设备从未修复，章节却写已修复）。"
           "该项需要模型判断；离线运行时系统会明确标注“需接入模型、不判定”，不会静默通过。")

    h2(doc, "第 5 步：编辑、锁定与局部重新即兴")
    p(doc, "在「结构树」选中一章 →「章节卡片」可直接编辑 →点「保存章节卡片编辑」。"
           "认可的情节可「锁定」（整章锁定或字段级锁定）。")
    p(doc, "想改某章时点「重新即兴选中章」：系统先给出候选（显示改动字段），你采纳后才写回；"
           "未选中的章节与锁定内容不受影响。重新即兴会保持该章的 POV 与参与人物不变（只改行动/冲突/结果等），"
           "保证人物记录与新章节一致。")

    h2(doc, "第 6 步：导出")
    p(doc, "点「导出结构 Markdown」得到结构文件；点「导出故事板（md/html）」得到可视化故事板"
           "（动机时间线 / 人物弧线 / 伏笔面板 / 偏离检测）。")

    h2(doc, "新建 / 打开项目")
    p(doc, "GUI 顶部有「项目」下拉与两个按钮：")
    bullets(doc, [
        "新建项目：填写名称与“世界观/大致剧情”，程序在 projects/ 下生成一套独立乐谱"
        "（story / characters / relationships / motifs / story_state），并自动打开。",
        "打开：从下拉选择已存在的项目，点“打开”；数据目录随之切换，规划/正文/历史都写进该项目"
        "（打开后会弹确认，右侧显示“项目：…”）。",
        "编辑/重建：修改项目简介并重建，程序按新简介重新生成整套设定；旧文件（含已保存结构）"
        "会移入 _backup_* 目录，之后需重新“生成创意→采纳并展开”。",
    ])
    p(doc, "命令行：python main.py project list / show <目录> / new \"名称\" --brief \"…\" / "
           "regen <目录> --file 简介.txt；之后用 --data projects/<目录> 继续。")

    # -- 4 ----------------------------------------------------------------
    h1(doc, "四、界面速览")
    h2(doc, "顶部工具栏")
    table(doc, ["控件", "说明"], [
        ["模型", "mock（离线演示）/ opencode（复用你已连接的模型）/ openai（兼容端点）"],
        ["评价", "文本相似度算法：lexical（默认、无依赖）/ embedding / hybrid；改后点“重载设定”"],
        ["重载设定", "使模型/评价模式/数据文件的改动生效"],
    ])
    h2(doc, "左侧控制区")
    table(doc, ["区域 / 按钮", "作用"], [
        ["场景 / 章节摘要框", "写故事前提；正文演练时每行一章"],
        ["即兴模式 / 视角 POV / Solo 色彩", "控制发散与文风"],
        ["结构模板", "线性因果 / 悬疑 / 多视角 / 群像 / 非线性拼图 / 自定义"],
        ["全书章数 / 采纳方案# / 融合方案A#、B#", "期望章数与候选编号"],
        ["结构参数 / 文本参数", "创意发散、因果变化、延迟揭露、连续性等"],
        ["① 结构即兴", "生成创意 / 融合 / 采纳展开 / 校验 / 卡片编辑 / 重新即兴 / 锁定 / 导出 / 偏离检测 / 故事板"],
        ["② 局部演练", "按结构渲染正文 / 试写一章 / 多版本探索 / 场景建议 / Jam Session"],
    ])
    h2(doc, "右侧标签页")
    table(doc, ["标签", "内容"], [
        ["创意候选", "多个故事方向"],
        ["全书结构", "梗概 + 分卷"],
        ["结构树 / 章节卡片", "逐章浏览与编辑"],
        ["卡片JSON", "整张章节卡片（含人物决策/认知/关系变化）的 JSON 编辑"],
        ["人物", "人物库、使用情况（参与章）、关系网络、闲置提示"],
        ["Jam", "四角色合奏产出的场景提案"],
        ["正文", "按结构渲染或试写的正文"],
        ["即兴决策 / 评价 / 建议", "本轮 Cue / 文本评价 / 下一轮建议"],
        ["状态 / 历史", "乐谱与故事状态 / 运行记录与趋势"],
    ])

    # -- 5 ----------------------------------------------------------------
    h1(doc, "五、核心概念：乐谱 · 结构 · 人物与关系")
    p(doc, "把设定想成爵士乐队的“乐谱”，AI 是与你合奏的乐手。", bold=True)
    h2(doc, "乐谱（data/ 目录）")
    table(doc, ["文件", "内容"], [
        ["story.yaml", "主题命题、世界硬约束/禁忌、情绪曲线、即兴参数"],
        ["characters.yaml", "人物库：id/层级/是否可作 POV + 叙事驱动字段"],
        ["relationships.yaml", "人物关系网络：需要 / 冲突 / 不对称 / 压力点 / 可能走向"],
        ["motifs.yaml", "意象定义（运行时意义演变写入 motif_state.yaml）"],
        ["story_state.yaml", "故事圣经：事实 / 怀疑 / 秘密 / 未决线索 / 近期事件"],
    ])
    h2(doc, "结构三层")
    bullets(doc, [
        "BookPlan（全书）：梗概、命题、冲突、关键转折、高潮、结局、分卷、人物 ids",
        "VolumePlan（分卷）：本卷目标、开端、高潮、阶段结果",
        "ChapterPlan（章节卡片）：功能、POV、起因、行动、冲突、结果、信息变化、人物变化、线索、意象、前后钩子，以及人物参与（参与人物 / 关键决定 / 认知变化）",
    ])
    h2(doc, "人物与关系网络")
    p(doc, "人物数量不设上限，按层级管理（示例给了 5 人）：")
    bullets(doc, [
        "core 核心（2–4）：决定主线走向",
        "supporting 重要配角（3–6）：握信息、代表不同利益、或有支线",
        "functional 功能性人物（按需）：证人、办事人员等，不必都有完整弧线",
    ])
    p(doc, "关键人物字段：id（稳定标识）、tier（层级）、pov（能否当视角）、"
           "desire / fear（欲望与恐惧）、blind_spots（认知局限，会在何处误判）、"
           "secrets（隐瞒）、capabilities / constraints（能做 / 受限）、current_goal（当下目标）。")
    p(doc, "关系网络比继续给人加形容词更重要：它决定“谁必须合作、谁可能互相阻碍、一个人的选择如何代价化另一个人”。")
    bullets(doc, [
        "每章可记录关系变化（谁与谁 / 如何变 / 因为什么）；转折类章节会依据关系里的“压力点/可能走向”自动生成。",
        "校验会提示：未定义的关系发生了变化、关系变化缺少原因、已定义的关系在全书中从未发生变化，以及“与主角隔离的人物关系组”。",
        "用 python main.py chars 或 GUI「人物」页可查看每条关系在哪些章发生变化。",
        "引用人物可用稳定 id、姓名或别名（系统会解析到同一人物）；主角由“核心且可作视角”的第一人确定，不依赖文件顺序。",
        "每章的人物决策/认知变化/关系变化是结构化数据，可在 GUI「卡片JSON」标签页直接编辑并保存。",
        "认知连续性：每章「认知变化」必须给出信息来源，来源不能指向不早于本章的时间；涉及作者保留的秘密时系统会提示确认揭露时机。",
        "结构化编辑：点「编辑人物/关系变化（表单）」用简表编辑本章的决策/认知/关系变化；也可在「卡片JSON」里直接改整张卡片。",
    ])
    h2(doc, "硬约束 vs 软约束")
    bullets(doc, [
        "硬约束（不可违反）：世界规则、禁忌、已确定事实、人物认知、未决线索的揭露时机",
        "软约束（可变化）：描述方式、场景顺序、对话风格、语言距离",
    ])

    # -- 6 ----------------------------------------------------------------
    h1(doc, "六、结构工作流详解")
    numbers(doc, [
        "输入前提，设定模式/模板/章数。",
        "生成创意方案（故事级发散；会把人物的欲望/盲点与关系网络一并纳入）。",
        "比较方案（因果机制/真相/关系/结局/差异度），可融合；前提改了会提示方案过期。",
        "采纳并展开：生成全书梗概（开端/触发事件/关键转折/高潮/结局处理）与逐章卡片；人物按层级分配到各章。",
        "校验：结构硬校验（缺章、编号、断链、未回收线索、闲置人物）+ 叙事审阅（含依据与章号）。",
        "编辑与锁定：卡片可编辑；可整章锁定或字段级锁定。",
        "局部重新即兴：先出候选并显示“改动字段”，采纳后才写回；未选中章节与锁定内容保持。",
        "导出：结构（Markdown）与故事板（Markdown/HTML）。",
    ])
    p(doc, "创意阶段提供两种对照：“五维并排对照”（因果机制 / 信息结构·真相 / 人物关系 / 结局 / 主题解释）"
           "与“融合对照”（融合后每个维度取自 A 还是 B）。结构评价在六维之外新增“关系演变”维度。")

    # -- 7 ----------------------------------------------------------------
    h1(doc, "七、局部演练：正文 / 多版本 / Jam")
    h2(doc, "按结构渲染正文")
    p(doc, "用已生成的章节摘要作为场景，顺序写成正文（可选）。按钮：「按结构渲染正文」。")
    p(doc, "在线正文预览：每次渲染都会生成一个自包含阅读页，并默认同步到远端——仓库根目录 "
           "preview/<项目>.html 与索引 preview/index.html。启用 GitHub Pages 后即可在 "
           "https://<用户名>.github.io/<仓库>/preview/ 阅读。"
           "开启方法：打开仓库的 Settings → Pages（地址形如 "
           "https://github.com/<用户名>/<仓库>/settings/pages，注意不是右上角头像里的账号设置），"
           "Build and deployment → Source 选 Deploy from a branch → Branch main、Folder /(root) → Save"
           "（预览在根目录 preview/，用默认根目录源即可；若选 /docs 源则地址变为 /docs/preview/）。"
           "不想自动推送可设环境变量 NJE_PREVIEW_AUTOPUSH=0，或命令加 --no-publish。")
    h2(doc, "试写与多版本")
    bullets(doc, [
        "「演奏一章（试写）」：局部即兴一段正文，含“即兴决策 Cue”“评价”。",
        "「多版本 · 探索路径」：同一主题多次独立即兴，得到不同叙事路径。",
    ])
    h2(doc, "Jam Session")
    p(doc, "点「Jam Session（导演/独奏/观察/编辑）」：四角色分工协作，"
           "Editor 会对候选打分排序、剔除与已知事实冲突者，整合成一个场景提案（不自动升级为正式设定）。")
    p(doc, "多轮合奏（rounds≥2）：观察者会先**反驳**独奏者（太笼统/直说主题/复述既有冲突），"
           "独奏者据此**修订**，Editor 再**归因**（哪个想法来自谁）。命令行：`jam --rounds 2`。")
    p(doc, "「导出故事板」产出的是**可交互 HTML**：可按视角筛选、搜索关键词、点击章节查看卡片；"
           "向量记忆用 `python main.py memory --rebuild` / `--query` 做章节语义检索。")

    # -- 8 ----------------------------------------------------------------
    h1(doc, "八、接入真实模型")
    p(doc, "默认 mock 只用于离线演示。要获得真正的创意与文笔，请接入真实模型（二者兼容）。")
    h2(doc, "方式一：opencode（复用你已连接的模型）")
    numbers(doc, [
        "另开一个终端启动服务器：opencode serve --port 4096",
        "在 JazzWriter 顶部把“模型”下拉选 opencode，点“重载设定”。",
        "可选环境变量：NJE_OPENCODE_URL（默认 http://127.0.0.1:4096）、NJE_OPENCODE_MODEL（如 anthropic/claude-…）。",
    ])
    h2(doc, "方式二：OpenAI 兼容端点（OpenAI / DeepSeek / Qwen…）")
    p(doc, "在启动 GUI 之前设置环境变量：")
    p(doc, 'pip install openai')
    p(doc, '$env:OPENAI_API_KEY = "你的key"')
    p(doc, '$env:NJE_BASE_URL = "https://api.deepseek.com/v1"')
    p(doc, '$env:NJE_MODEL = "deepseek-chat"')
    p(doc, "然后顶部“模型”选 openai，点“重载设定”。")
    p(doc, "提示：结构在离线（mock）下用确定性模板，会被标记为 offline_template；只有接入模型才是模型创作。")

    # -- 9 ----------------------------------------------------------------
    h1(doc, "九、命令行速查")
    table(doc, ["命令", "作用"], [
        ["python main.py status", "查看主题/人物/和声/动机/故事状态"],
        ["python main.py chars", "查看人物库、使用情况与关系网络"],
        ["python main.py memory --rebuild", "从当前结构重建向量记忆（语义检索库）"],
        ["python main.py memory --query \"...\"", "按语义检索最相近的章节"],
        ["python main.py project list", "列出已有项目（projects/）"],
        ["python main.py project new \"名称\" --brief \"…\"", "按世界观/剧情生成新项目"],
        ["python main.py outline \"前提\" --branches 3", "发散故事方案"],
        ["python main.py outline \"前提\" --fuse 1 2", "融合两个方案"],
        ["python main.py outline \"前提\" --build --pick 2 --chapters 12 --template mystery", "采纳并展开结构"],
        ["python main.py outline --validate", "结构校验"],
        ["python main.py outline --lock 3", "锁定第 3 章"],
        ["python main.py outline --lock-field 3 conflict", "锁定第 3 章的 conflict 字段"],
        ["python main.py outline --reimprovise 5", "生成第 5 章候选（不覆盖）"],
        ["python main.py outline --reimprovise 5 --adopt 1 --regen all", "采纳候选并重生成后续"],
        ["python main.py outline --edit-json 3 --file card.json", "用 JSON 更新第 3 章（支持嵌套人物/关系字段）"],
        ["python main.py outline --export", "导出结构 Markdown"],
        ["python main.py book --from-plan", "按结构渲染正文"],
        ["python main.py jam \"场景\" --pov 林 --render", "Jam 合奏并渲染"],
        ["python main.py report --open", "偏离检测 + 导出故事板并打开"],
        ["python main.py history", "运行历史与趋势"],
        ["python main.py advice --apply", "应用参数建议"],
        ["python main.py init", "写入示例乐谱（含人物与关系）"],
    ])

    # -- 10 ---------------------------------------------------------------
    h1(doc, "十、常见问题（FAQ）")
    qa = [
        ("生成的内容很模板化？",
         "那是在用离线 mock。请接入真实模型（opencode 或 openai），结构与文笔主要依赖模型。"),
        ("需要联网 / API Key 吗？",
         "离线 mock 不需要；真实模型需要。"),
        ("人物只能有几个？能加吗？",
         "不设上限。直接编辑 data/characters.yaml 增加人物即可；用 tier 分层、pov 决定能否当视角、"
         "约定 id 便于引用。短篇 2–3 人即可，群像可更多。"),
        ("关系网络有什么用？",
         "它让人物之间产生真实的合作与冲突，是故事动力来源。编辑 data/relationships.yaml 即可。"
         "用 python main.py chars 或 GUI「人物」页查看。"),
        ("我的修改会被覆盖吗？",
         "不会。锁定的章节与字段不会被改写；“重新即兴”先给候选，采纳后才写回。"),
        ("如何得到小说正文？",
         "先有结构，再用「按结构渲染正文」或 python main.py book --from-plan。"),
        ("能复现同一次结果吗？",
         "固定“随机种子 seed”可复现即兴决策；正文由模型采样，可能略有差异。"),
        ("支持哪些模型？",
         "opencode（复用已连接模型）与任意 OpenAI 兼容端点。"),
        ("中文乱码？",
         "GUI 无此问题；命令行若乱码是终端编码，改用 Windows Terminal 或设置 UTF-8。"),
        ("看到“人物未参与任何章节”？",
         "这是人物使用检查：登记了却没在结构里出现。可增加其戏份，或降低其层级/暂不使用。"),
    ]
    for q, a in qa:
        p(doc, "问：" + q, bold=True)
        p(doc, "答：" + a)

    # -- 11 ---------------------------------------------------------------
    h1(doc, "十一、文件与目录说明")
    table(doc, ["路径", "说明"], [
        ["data/", "乐谱：主题、人物、关系、意象、故事状态（可编辑）"],
        ["output/chapters/", "生成的章节正文（.md / .json）"],
        ["output/runs/", "每次运行的完整记录（cue/prompt/评价）"],
        ["<数据目录>/book_plan.json", "当前全书结构（自动保存）"],
        ["<数据目录>/book_outline.md", "结构导出"],
        ["<数据目录>/story_report.md / .html", "故事板导出（交互式 HTML）"],
        ["<数据目录>/vector_store.json", "向量记忆（语义检索库）"],
        ["data/motif_state.yaml / jazz_state.yaml", "运行时状态（本地，不入库）"],
    ])
    p(doc, "提示：用 --data <目录> 切换独立小说项目时，全书规划、结构导出、向量库、故事板、运行记录、"
           "章节正文与参数状态都保存在该项目目录下，不同项目互不干扰。")

    # -- 12 ---------------------------------------------------------------
    h1(doc, "十二、术语表")
    table(doc, ["术语", "含义"], [
        ["Theme 主题", "全书核心命题，只被暗示、不直说"],
        ["Harmony 和声", "世界规则、禁忌、情绪曲线等约束"],
        ["Motif 动机", "反复出现且意义不断变化的意象"],
        ["Character 人物", "含 id/层级/视角资格与叙事驱动字段（欲望/恐惧/盲点/秘密/能力/限制/目标）"],
        ["Relationship 关系", "人物之间的需要、冲突、不对称与可能走向"],
        ["tier 层级", "core 核心 / supporting 重要配角 / functional 功能性"],
        ["Cue 乐谱", "某一轮即兴的决策（色彩/操作/意外/节奏）"],
        ["Take", "一次生成的文本版本"],
        ["BookPlan / VolumePlan / ChapterPlan", "全书 / 分卷 / 章节卡片三层结构"],
        ["Jam Session", "导演→独奏→观察→编辑 顺序合奏"],
        ["offline_template", "离线模板生成的内容（非模型创作）"],
    ])

    return doc


def main() -> int:
    doc = build()
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"written: {OUT}  ({len(doc.paragraphs)} paragraphs)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
